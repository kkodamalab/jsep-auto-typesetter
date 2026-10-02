import shutil
import subprocess
from pathlib import Path

import pytest

from backend.converter import ConversionError, extract_docx
from backend.models import PdfRequest
from backend.pdf import generate_pdf


def _require_toolchain():
    missing = [tool for tool in ("pandoc", "lualatex", "pdfinfo", "pdftotext", "pdfimages", "pdftoppm")
               if not shutil.which(tool)]
    if missing:
        pytest.fail(f"Integration toolchain is incomplete: {', '.join(missing)}")


def _fictional_docx(path: Path):
    import docx
    from PIL import Image
    drawing = path.parent / "架空グラフ.png"
    canvas = Image.new("RGB", (480, 240), "white")
    for x in range(40, 440):
        canvas.putpixel((x, 200 - ((x - 40) // 3) % 150), (20, 83, 65))
    canvas.save(drawing)

    document = docx.Document()
    document.core_properties.title = '架空: 「教育&学習」_100% #1'
    document.core_properties.author = "山田 花子 & O'Connor"
    document.add_heading("抄録", level=1)
    document.add_paragraph("これは変換試験専用の架空の日本語抄録である。")
    document.add_heading("はじめに", level=1)
    document.add_paragraph("日本語本文と特殊文字 # $ % & _ { } ~ ^ \\ を安全に変換する。")
    document.add_picture(str(drawing), width=docx.shared.Cm(10))
    document.add_paragraph("図1　架空データの推移", style="Caption")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "条件", "値"
    table.cell(1, 0).text, table.cell(1, 1).text = "架空条件", "42"
    document.add_paragraph("表1　架空実験の結果", style="Caption")
    paragraph = document.add_paragraph("質量とエネルギーの関係: ")
    from docx.oxml import parse_xml
    paragraph._p.append(parse_xml(
        '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
        '<m:r><m:t>E=mc²</m:t></m:r></m:oMath>'))
    document.add_heading("参考文献", level=1)
    document.add_paragraph("山田花子（2026）架空研究の方法。架空出版。")
    document.save(path)


@pytest.mark.integration
def test_fictional_docx_to_verified_pdf(tmp_path):
    _require_toolchain()
    source = tmp_path / "fictional.docx"
    _fictional_docx(source)
    artifacts = Path("test-artifacts")
    artifacts.mkdir(exist_ok=True)
    shutil.copy2(source, artifacts / "fictional-manuscript.docx")
    manuscript = extract_docx(source, tmp_path)
    expected_title = '架空: 「教育&学習」_100% #1'
    expected_author = "山田 花子 & O'Connor"
    assert manuscript.title == expected_title
    assert [author.name for author in manuscript.authors] == [expected_author]
    assert "media/" in manuscript.body_markdown
    assert "図1" in manuscript.body_markdown and "表1" in manuscript.body_markdown

    # Exercise edited metadata containing YAML/LaTeX-special characters.
    manuscript.title += ': 改訂版 "安全性"'
    try:
        output = generate_pdf(PdfRequest(manuscript_id="integration", manuscript=manuscript), tmp_path)
    except ConversionError:
        # CI processes fictional data only. Preserve compiler diagnostics without
        # enabling this export path for real uploads in the runtime application.
        for name in ("manuscript.tex", "manuscript.log"):
            diagnostic = tmp_path / name
            if diagnostic.exists():
                shutil.copy2(diagnostic, artifacts / name)
        raise
    assert output.read_bytes().startswith(b"%PDF-") and output.stat().st_size > 10_000

    shutil.copy2(output, artifacts / "fictional-manuscript.pdf")
    subprocess.run(["pdftoppm", "-png", "-r", "120", str(output), str(artifacts / "page")], check=True)

    info = subprocess.run(["pdfinfo", str(output)], check=True, capture_output=True, text=True).stdout
    pages = int(next(line.split(":", 1)[1] for line in info.splitlines() if line.startswith("Pages:")))
    assert pages >= 1
    text = subprocess.run(["pdftotext", str(output), "-"], check=True, capture_output=True, text=True).stdout
    for expected in ("教育&学習", "山田 花子", "O'Connor", "日本語本文", "架空条件", "E=mc", "参考文献"):
        assert expected in text, f"PDF text is missing or garbled: {expected}"
    images = subprocess.run(["pdfimages", "-list", str(output)], check=True, capture_output=True, text=True).stdout
    assert len(images.splitlines()) > 2, "PDF contains no embedded image"

    assert len(list(artifacts.glob("page-*.png"))) == pages
